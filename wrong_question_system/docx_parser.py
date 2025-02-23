# docx_parser.py
import os
import re
import docx
from log_config import get_logger

import json
with open("config.json", "r", encoding="utf-8") as f:
    config = json.load(f)

log_path = config.get("log_path")
logger = get_logger(__name__, log_file=log_path)

def extract_text_and_images_from_docx(docx_file, base_image_dir):
    try:
        base = os.path.splitext(os.path.basename(docx_file))[0]
        image_dir = os.path.join(base_image_dir, base)
        if not os.path.exists(image_dir):
            os.makedirs(image_dir)
            logger.info("创建图片保存目录: %s", image_dir)
        
        doc_object = docx.Document(docx_file)
        image_counter = 0
        final_lines = []
        
        for paragraph in doc_object.paragraphs:
            paragraph_chunks = []
            for run in paragraph.runs:
                drawing_el = run._element.xpath('.//*[local-name()="drawing"]')
                if drawing_el:
                    blip_elems = run._element.xpath('.//*[local-name()="blip"]')
                    if blip_elems:
                        blip = blip_elems[0]
                        rid = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                        if rid:
                            image_counter += 1
                            part = doc_object.part.related_parts[rid]
                            image_filename = f"image{image_counter}.png"
                            out_path = os.path.join(image_dir, image_filename)
                            with open(out_path, "wb") as f:
                                f.write(part._blob)
                            logger.info("提取图片: %s", out_path)
                            paragraph_chunks.append(f"[{image_filename}]")
                        else:
                            paragraph_chunks.append("[image_no_rid]")
                    else:
                        paragraph_chunks.append("[image_no_blip]")
                else:
                    paragraph_chunks.append(run.text)
            line_text = "".join(paragraph_chunks).strip()
            if line_text:
                final_lines.append(line_text)
        return "\n".join(final_lines), image_dir
    except Exception as e:
        logger.error("提取 docx 文本和图片时出错: %s", e, exc_info=True)
        raise

def replace_image_placeholders_with_captions(text, image_dir, processor, model, device):
    # 注意这里导入 image_caption 模块中的函数
    from image_caption import chinese_image_caption_inference
    pattern = re.compile(r'\[image(\d+)\.png\]', re.IGNORECASE)
    
    def _replace_func(match):
        image_number = match.group(1)
        image_filename = f"image{image_number}.png"
        image_path = os.path.join(image_dir, image_filename)
        logger.info("正在处理图片: %s", image_path)
        if os.path.exists(image_path):
            caption = chinese_image_caption_inference(image_path, processor, model, device)
            return f"[{image_path}] {caption} [/]"
        else:
            logger.warning("图片未找到: %s", image_path)
            return f"[图片 {image_filename} 未找到] [/]"
    try:
        new_text = pattern.sub(_replace_func, text)
        return new_text
    except Exception as e:
        logger.error("替换图片占位符时出错: %s", e, exc_info=True)
        return text

def remove_line_numbering(text: str) -> str:
    import re
    return re.sub(r'\d+\.\s+', '', text)
